"""레거시 명세 자산 생성기 — 단일 소스 catalog.CATALOG 로부터 만든다.

  ctlg.swagger2.json            CTLG Swagger 2.0 (세션 쿠키)
  dhgw.openapi3.yaml            DHGW OpenAPI 3.0 (OAuth2 client_credentials)
  finl_인터페이스정의서.xlsx     FINL 엑셀 인터페이스정의서 (기계가독 명세 없음)

[매설 결함 — 제거 금지] FINL 정의서에는 doc_only 필드(ACNT_DIV_CD)를 응답 항목으로 적는다. 실제 서버는 돌려주지 않는다.
실행: python assets/make_assets.py   (apps/legacy-pps 에서)
"""
import json
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from legacy_pps.common.catalog import CATALOG  # noqa: E402

DESC = {
    # CTLG
    "PRDCT_CLSFC_NO": "물품분류번호 8자리(대·중·소·세분류 각 2자리, UNSPSC 기반)", "DTL_PRDCT_CLSFC_NO": "세부품명번호 10자리(물품분류번호+2자리)",
    "DTL_PRDCT_CLSFC_NM": "세부품명", "SGMNT_CD": "대분류(Segment) 2자리", "FMLY_CD": "중분류(Family) 4자리", "CLS_CD": "소분류(Class) 6자리",
    "UNIT_NM": "단위", "PRDCT_IDNT_NO": "물품식별번호 8자리", "MODEL_NM": "모델명", "MAKER_NM": "제조사명",
    "MAKER_BIZRNO": "제조사 사업자등록번호", "UNIT_PRCE": "단가(원)", "BIZRNO": "사업자등록번호", "CORP_NM": "업체명",
    "RPRSV_NM": "대표자명(가림)", "RGN_NM": "소재지 시도", "ENTRPS_DIV_NM": "업체구분", "TELNO": "전화번호(가림)",
    # DHGW
    "bidNtceNo": "입찰공고번호", "bidNtceOrd": "입찰공고차수", "bidNtceNm": "입찰공고명", "ntceKindNm": "공고종류명",
    "ntceInsttCd": "공고기관코드", "ntceInsttNm": "공고기관명", "dminsttCd": "수요기관코드", "dminsttNm": "수요기관명",
    "bidMethdNm": "입찰방식명", "cntrctCnclsMthdNm": "계약체결방법명", "bidNtceDt": "입찰공고일시 YYYYMMDDHHMM",
    "bidBeginDt": "입찰개시일시", "bidClseDt": "입찰마감일시", "opengDt": "개찰일시", "presmptPrce": "추정가격(원)",
    "asignBdgtAmt": "배정예산금액(원)", "prdctClsfcNo": "물품분류번호", "prdctClsfcNoNm": "품명", "prdctQty": "수량", "prdctUnit": "단위",
    "prtcptCnum": "참가업체수", "bidwinnrNm": "최종낙찰업체명", "bidwinnrBizno": "최종낙찰업체 사업자등록번호",
    "sucsfbidAmt": "최종낙찰금액(원)", "sucsfbidRate": "최종낙찰률(%)", "progrsDivCdNm": "진행구분명",
    "untyCntrctNo": "통합계약번호", "cntrctChgOrd": "계약변경차수", "cntrctNm": "계약명", "cntrctCnclsDate": "계약체결일자 YYYYMMDD",
    "thtmCntrctAmt": "금차계약금액(원)", "cntrctInsttNm": "계약기관명", "corpNm": "업체명", "bizno": "사업자등록번호",
    "ntceNo": "입찰공고번호", "dlvrTmlmtDate": "납품기한일자", "dlvrReqNo": "납품요구번호", "dlvrReqRcptDate": "납품요구접수일자",
    "prdctIdntNo": "물품식별번호", "prdctIdntNoNm": "품명", "modelNm": "모델명", "dlvrReqQty": "납품요구수량",
    "prdctUprc": "단가(원)", "dlvrReqAmt": "납품요구금액(원)", "dlvrReqSttsCd": "납품요구상태코드 10접수 20납품중 30검수완료",
    # FINL
    "PYMNT_REQ_NO": "대금청구번호", "CNTRCT_NO": "계약번호 또는 납품요구번호", "DMINSTT_CD": "수요기관코드", "REQ_YMD": "청구일자",
    "REQ_AMT": "청구금액(원, 천단위 쉼표)", "PYMNT_STTS_CD": "지급상태코드 10청구접수 20검사완료 30지급완료 40반려",
    "PYMNT_STTS_NM": "지급상태명", "PYMNT_YMD": "지급일자", "BDGT_ACNT_NM": "예산회계명", "BDGT_YR": "예산연도",
    "ASIGN_AMT": "배정액(원)", "EXCUT_AMT": "집행액(원)", "RMND_AMT": "잔액(원)",
}


def _params2(op):
    return [{"name": n, "in": "query" if op["method"] == "GET" else "formData", "type": "string", "required": req, "description": d}
            for n, d, req in op["params"]]


def swagger2(sys_key, base, title):
    paths = {}
    for op in CATALOG[sys_key]["ops"]:
        props = {dst: {"type": "string", "description": DESC.get(dst, dst)} for _, dst in op["fields"]}
        paths[op["path"]] = {op["method"].lower(): {
            "operationId": op["op"], "summary": op["summary"], "produces": ["application/json"], "parameters": _params2(op),
            "responses": {"200": {"description": "정상(세션 만료 시에도 200 + text/html)", "schema": {"type": "object", "properties": {
                "resultCode": {"type": "string"}, "resultMsg": {"type": "string"}, "totCnt": {"type": "integer"},
                "resultList": {"type": "array", "items": {"type": "object", "properties": props}}}}}},
            "security": [{"JSESSIONID": []}]}}
    return {"swagger": "2.0", "info": {"title": title, "version": "1.3.2", "description": "시연용 가상 시스템. actionLogin.do 로 세션을 받는다."},
            "host": "10.70.1.10:8001", "basePath": "/", "schemes": ["http"],
            "securityDefinitions": {"JSESSIONID": {"type": "apiKey", "in": "header", "name": "Cookie"}}, "paths": paths}


def openapi3(sys_key, title):
    paths = {}
    for op in CATALOG[sys_key]["ops"]:
        props = {dst: {"type": "string", "description": DESC.get(dst, dst)} for _, dst in op["fields"]}
        paths[op["path"]] = {op["method"].lower(): {
            "operationId": op["op"], "summary": op["summary"],
            "parameters": [{"name": n, "in": "query", "required": req, "description": d, "schema": {"type": "string"}} for n, d, req in op["params"]],
            "responses": {"200": {"description": "정상", "content": {"application/json": {"schema": {"type": "object", "properties": {
                "response": {"type": "object", "properties": {
                    "header": {"type": "object", "properties": {"resultCode": {"type": "string"}, "resultMsg": {"type": "string"}}},
                    "body": {"type": "object", "properties": {"totalCount": {"type": "integer"}, "pageNo": {"type": "integer"},
                                                              "numOfRows": {"type": "integer"},
                                                              "items": {"type": "array", "items": {"type": "object", "properties": props}}}}}}}}}}},
                          "401": {"description": "토큰 없음·만료 {\"error\":\"invalid_token\"}"}}}}
    return {"openapi": "3.0.3", "info": {"title": title, "version": "2.1.0", "description": "시연용 가상 시스템. 값은 모두 문자열이다."},
            "servers": [{"url": "http://10.70.1.10:8002"}],
            "components": {"securitySchemes": {"oauth2": {"type": "oauth2", "flows": {"clientCredentials": {"tokenUrl": "/oauth2/token", "scopes": {}}}}}},
            "security": [{"oauth2": []}], "paths": paths}


def finl_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "인터페이스 목록"
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="D9E1F2")
    ws.append(["FINL 재정 연계 인터페이스 정의서 (시연용 가상 시스템)"])
    ws.append(["문서번호", "FINL-IF-2015-003", "개정", "v1.4 (2019.03)", "작성", "재정연계 운영팀"])
    ws.append([])
    ws.append(["IF ID", "인터페이스명", "방식", "URL", "인증"])
    for c in ws[4]:
        c.font, c.fill, c.border = Font(bold=True), head, box
    for i, op in enumerate(CATALOG["finl"]["ops"], 1):
        ws.append([f"IF-FINL-{i:03d}", op["summary"], op["method"], op["path"], "HTTP 헤더 X-API-KEY"])
    for i, op in enumerate(CATALOG["finl"]["ops"], 1):
        s = wb.create_sheet(f"IF-FINL-{i:03d}")
        s.append([f"IF-FINL-{i:03d}  {op['summary']}  ({op['op']})"])
        s.append(["구분", "항목ID", "항목명", "타입", "필수", "비고"])
        for c in s[2]:
            c.font, c.fill, c.border = Font(bold=True), head, box
        for n, d, req in op["params"]:
            s.append(["요청", n, d, "VARCHAR(20)", "Y" if req else "N", ""])
        for _, dst in op["fields"]:
            s.append(["응답", dst, DESC.get(dst, dst), "VARCHAR(20)", "Y", "금액은 천단위 쉼표" if dst.endswith("_AMT") else ""])
        for n, typ, req, d in op.get("doc_only", []):      # 매설 결함 — 정의서에만 존재
            s.append(["응답", n, d, typ, req, ""])
        for row in s.iter_rows(min_row=3):
            for c in row:
                c.border, c.alignment = box, Alignment(vertical="top", wrap_text=True)
        for col, w in zip("ABCDEF", (8, 18, 40, 14, 6, 22)):
            s.column_dimensions[col].width = w
    wb.save(path)


if __name__ == "__main__":
    (HERE / "ctlg.swagger2.json").write_text(json.dumps(swagger2("ctlg", "", CATALOG["ctlg"]["title"]), ensure_ascii=False, indent=2), encoding="utf-8")
    (HERE / "dhgw.openapi3.yaml").write_text(yaml.safe_dump(openapi3("dhgw", CATALOG["dhgw"]["title"]), allow_unicode=True, sort_keys=False), encoding="utf-8")
    finl_xlsx(HERE / "finl_인터페이스정의서.xlsx")
    print("assets ok")
